import argparse
import json
import os
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


GREEN = "E2F0D9"
RED = "FCE4D6"
GREY = "F2F2F2"


def replace_paragraph_text(paragraph, text):
    """Replace text while retaining the formatting of the first run."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_cell_text(cell, text, bold=None, color=None, size=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def clear_table_body(table):
    """Keep the header row and remove all placeholder body rows."""
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)


def mark_header_repeat(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def add_hyperlink(paragraph, text, target):
    """Add a clickable external or relative file link to a paragraph."""
    relationship = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(colour)
    properties.append(underline)
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def report_relative_path(artifact_path, report_directory):
    """Create a portable link from the report folder to its evidence file."""
    artifact = Path(artifact_path)
    if not artifact.is_absolute():
        artifact = (Path.cwd() / artifact).resolve()
    return Path(os.path.relpath(artifact, report_directory.resolve())).as_posix()


def enable_field_updates(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def make_variable_toc_page_safe(doc):
    """Avoid retaining the template's old final-page number before Word updates fields."""
    for hyperlink in doc._element.xpath(".//w:hyperlink"):
        text_nodes = hyperlink.xpath(".//w:t")
        combined = "".join(node.text or "" for node in text_nodes)
        if combined.startswith("About Metricell"):
            for node in text_nodes:
                if (node.text or "").strip().isdigit():
                    node.text = "Final page"


def clean_failure(text):
    text = re.sub(r"\x1b\[[0-9;]*m", "", text or "")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    useful = [
        line for line in lines
        if "AssertionError" in line or "Error:" in line or "Timeout" in line
    ]
    return (useful[-1] if useful else (lines[-1] if lines else "No failure detail was recorded."))[:700]


def base_test_name(name):
    """Remove pytest browser/parameter suffixes such as [chromium]."""
    return (name or "").split("[", 1)[0]


def actual_result(result, meta):
    recorded_actual = result.get("actual_result")
    if recorded_actual:
        return recorded_actual
    if result["outcome"] == "passed":
        return meta.get(
            "success_actual",
            "The automated checks completed and the expected behaviour was observed.",
        )
    if result["outcome"] == "skipped":
        return "Test was skipped during this execution."
    failure = clean_failure(result.get("failure_summary") or result.get("failure"))
    location = result.get("current_url")
    if location:
        return f"Test failed at {location}. Observed result: {failure}"
    return f"Test failed. Observed result: {failure}"


def build_report(
    results_path="test-results/results.json",
    cases_path="test_cases.json",
    template_path="Test Document Template.docx",
    output_dir="reports",
    environment="https://smartnetworkv4-o2-uk-dev.metricell.com/",
    ticket_key=None,
):
    results = json.loads(Path(results_path).read_text(encoding="utf-8"))
    cases = json.loads(Path(cases_path).read_text(encoding="utf-8"))
    generated_at = datetime.fromisoformat(
        results["generated_at"].replace("Z", "+00:00")
    ).astimezone()
    tests = results.get("tests", [])
    if ticket_key:
        tests = [
            test for test in tests
            if ticket_key in cases.get(base_test_name(test["name"]), {}).get("jira_keys", [])
        ]
    suite_version = results.get("suite_version", "Metricell V4 test suite")
    execution_mode = results.get("execution_mode", "Automated")
    suite_number = suite_version.rsplit(" ", 1)[-1]

    if not tests:
        raise ValueError("No test results were found")

    doc = Document(template_path)
    enable_field_updates(doc)
    make_variable_toc_page_safe(doc)

    # Cover-page title and subtitle.
    titled = [p for p in doc.paragraphs if p.text.strip() == "Test Example"]
    if len(titled) >= 2:
        replace_paragraph_text(titled[0], "Automated QA Execution Report")
        subtitle = "Metricell Smart Network V4 Functional Regression Suite"
        if ticket_key:
            subtitle = f"Metricell Smart Network V4 Evidence Pack — {ticket_key}"
        replace_paragraph_text(titled[1], subtitle)

    # Cover-page metadata.
    cover = doc.tables[0]
    cover_values = [
        "Company Confidential",
        "Metricell Smart Network V4 Automated QA Execution",
        suite_number,
        generated_at.strftime("%d/%m/%Y"),
        "Automated Test Suite",
        "Quality Assurance/Metricell",
    ]
    for row, value in zip(cover.rows, cover_values):
        set_cell_text(row.cells[1], value, color="FFFFFF")

    # Revision history.
    revision = doc.tables[1]
    set_cell_text(revision.rows[1].cells[0], suite_number)
    set_cell_text(revision.rows[1].cells[1], generated_at.strftime("%d/%m/%Y"))
    set_cell_text(revision.rows[1].cells[2], "Automated Test Suite")
    set_cell_text(revision.rows[1].cells[3], "Automated daily test execution report")
    for row in revision.rows[2:]:
        for cell in row.cells:
            set_cell_text(cell, "")

    # Test cycle information.
    cycle = doc.tables[2]
    browser = next(
        (t.get("browser_version") for t in tests if t.get("browser_version")),
        "Chromium via Playwright",
    )
    cycle_values = [
        "Automated",
        environment,
        f"Chromium {browser}" if browser != "Chromium via Playwright" else browser,
        generated_at.strftime("%d/%m/%Y %H:%M"),
        f"Test suite: {suite_version}. Execution mode: {execution_mode}. "
        "Automated browser checks only; no manual exploratory testing included.",
    ]
    for row, value in zip(cycle.rows, cycle_values):
        set_cell_text(row.cells[1], value)

    # Summary by functional area.
    grouped = defaultdict(list)
    for test in tests:
        grouped[cases.get(base_test_name(test["name"]), {}).get("area", "General")].append(test)

    summary = doc.tables[3]
    clear_table_body(summary)
    mark_header_repeat(summary.rows[0])
    for area in sorted(grouped):
        area_tests = grouped[area]
        failures = sum(t["outcome"] == "failed" for t in area_tests)
        row = summary.add_row()
        values = [area, "Y" if failures == 0 else "", "Y" if failures else "", failures]
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)
        shade_cell(row.cells[0], GREY)
        if failures:
            shade_cell(row.cells[2], RED)
            shade_cell(row.cells[3], RED)
        else:
            shade_cell(row.cells[1], GREEN)

    # Replace template-specific requirement wording with demo scope.
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Please note that all test cases"):
            replace_paragraph_text(
                paragraph,
                "This report records automated Metricell Smart Network V4 checks executed by Playwright. "
                "It covers authenticated navigation, safe feature checks, GIS functionality and API monitoring.",
            )

    preconditions = doc.tables[4]
    set_cell_text(preconditions.rows[0].cells[1], "Metricell development access; saved login reused or refreshed interactively")
    set_cell_text(preconditions.rows[0].cells[3], "Metricell V4 automated QA suite")
    set_cell_text(preconditions.rows[0].cells[5], f"Generated by {suite_version}")

    output = Path(output_dir)
    output.mkdir(exist_ok=True)

    # Detailed test-case rows.
    detailed = doc.tables[5]
    clear_table_body(detailed)
    mark_header_repeat(detailed.rows[0])

    sorted_tests = sorted(
        tests,
        key=lambda result: (
            cases.get(base_test_name(result["name"]), {}).get("area", "General"),
            cases.get(base_test_name(result["name"]), {}).get("id", "ZZZ"),
        ),
    )

    for position, result in enumerate(sorted_tests, start=1):
        meta = cases.get(base_test_name(result["name"]), {})
        row = detailed.add_row()
        values = [
            meta.get("id", f"TC-{position:03d}"),
            meta.get("description", result["name"].replace("_", " ").title()),
            meta.get("steps", "Run the automated Playwright scenario."),
            meta.get("expected", "Expected behaviour is observed."),
            actual_result(result, meta),
            "Failed" if result["outcome"] in {"failed", "error"} else result["outcome"].title(),
        ]
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value, size=8.5)
        row.cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_cell(row.cells[5], GREEN if result["outcome"] == "passed" else RED)

        evidence = []
        if result.get("video") and Path(result["video"]).exists():
            evidence.append(("Open test video", result["video"]))
        if result.get("visual_screenshot") and Path(result["visual_screenshot"]).exists():
            evidence.append(("Open final-state screenshot", result["visual_screenshot"]))
        if result.get("trace") and Path(result["trace"]).exists():
            evidence.append(("Open failure trace", result["trace"]))
        if result.get("visual_current") and Path(result["visual_current"]).exists():
            evidence.append(("Open visual comparison image", result["visual_current"]))
        if result.get("visual_difference") and Path(result["visual_difference"]).exists():
            evidence.append(("Open visual difference", result["visual_difference"]))

        if evidence:
            paragraph = row.cells[4].add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            label = paragraph.add_run("Evidence: ")
            label.bold = True
            label.font.size = Pt(8.5)

            for link_number, (link_text, artifact_path) in enumerate(evidence):
                if link_number:
                    paragraph.add_run(" | ")
                target = report_relative_path(artifact_path, output)
                add_hyperlink(paragraph, link_text, target)

        screenshot = result.get("screenshot")
        if screenshot and Path(screenshot).exists():
            evidence_row = detailed.add_row()
            merged = evidence_row.cells[0].merge(evidence_row.cells[-1])
            merged.text = ""
            paragraph = merged.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(screenshot, width=Inches(6.0))
            caption = merged.add_paragraph("Final-state screenshot captured automatically for this test")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if caption.runs:
                caption.runs[0].italic = True
                caption.runs[0].font.size = Pt(8)

    prefix = f"{ticket_key}_Evidence" if ticket_key else "Metricell_QA_Report"
    filename = output / f"{prefix}_{generated_at:%Y-%m-%d_%H%M%S}.docx"
    doc.save(filename)
    return filename


def build_ticket_reports(
    results_path="test-results/results.json",
    cases_path="test_cases.json",
    template_path="Test Document Template.docx",
    output_dir="reports/tickets",
):
    """Generate one Word evidence pack for every Jira key represented in results."""
    results = json.loads(Path(results_path).read_text(encoding="utf-8"))
    cases = json.loads(Path(cases_path).read_text(encoding="utf-8"))
    ticket_keys = sorted({
        key
        for result in results.get("tests", [])
        for key in cases.get(base_test_name(result["name"]), {}).get("jira_keys", [])
    })
    return [
        build_report(results_path, cases_path, template_path, output_dir, ticket_key=key)
        for key in ticket_keys
    ]


def main():
    parser = argparse.ArgumentParser(
        description="Populate the Metricell Word template from pytest results"
    )
    parser.add_argument("--results", default="test-results/results.json")
    parser.add_argument("--cases", default="test_cases.json")
    parser.add_argument("--template", default="Test Document Template.docx")
    parser.add_argument("--output-dir", default="reports")
    parser.add_argument("--ticket-reports", action="store_true", help="Also generate one evidence pack per Jira ticket")
    args = parser.parse_args()
    output = build_report(
        args.results, args.cases, args.template, args.output_dir
    )
    print(output.resolve())
    if args.ticket_reports:
        for ticket_report in build_ticket_reports(args.results, args.cases, args.template, Path(args.output_dir) / "tickets"):
            print(ticket_report.resolve())


if __name__ == "__main__":
    main()

