"""Regression checks for guide evidence identity, including false-positive images."""

import hashlib
import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from PIL import Image

sys.dont_write_bytecode = True
MODULE = Path(__file__).with_name("verify-user-guide.py")
SPEC = importlib.util.spec_from_file_location("verify_user_guide", MODULE)
VERIFY = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(VERIFY)


class FakePage:
    def extract_text(self):
        return "New feature: Open GIS"


class FakeReader:
    pages = [FakePage()]


class GuideVerificationTests(unittest.TestCase):
    def setUp(self):
        root = MODULE.parents[1] / ".guide-staging" / "test-tmp"
        root.mkdir(parents=True, exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(dir=root)
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.evidence = self.root / "evidence"
        self.ticket = "TEST2-99"
        self.sources = []
        for index, colour in enumerate(("navy", "orange"), 1):
            source = self.evidence / self.ticket / "screenshots" / f"criterion-{index}.png"
            source.parent.mkdir(parents=True, exist_ok=True)
            Image.new("RGB", (80, 40), colour).save(source)
            self.sources.append(source)
        self.update = self.root / "update.json"
        self.update.write_text(json.dumps({"ticket": self.ticket, "title": "Open GIS", "steps": ["Open GIS."],
                                           "screenshots": [f"screenshots/{path.name}" for path in self.sources]}), encoding="utf-8")
        self.docx = self.root / "guide.docx"
        self.pdf = self.root / "guide.pdf"
        self.pdf.write_bytes(b"nonempty mocked PDF")

    def save_docx(self, screenshots):
        document = Document()
        document.add_paragraph("[[AUTO_GUIDE_CONTENT]]")
        document.add_paragraph("[[AUTO_GUIDE_UPDATE:TEST2-99]]")
        document.add_paragraph("Open GIS")
        document.add_paragraph("Open GIS.")
        for source in screenshots:
            document.add_picture(str(source))
        document.save(self.docx)

    def observed(self, sources):
        result = []
        for source in sources:
            with Image.open(source) as image:
                pixels = image.convert("RGB")
            result.append({"page": 1, "name": source.name, "sourceSha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                           "visualFingerprint": VERIFY.REPORT_VERIFIER.image_fingerprint(pixels), "image": pixels})
        return result

    def check(self, observed):
        with patch.object(VERIFY, "PdfReader", return_value=FakeReader()), patch.object(VERIFY.REPORT_VERIFIER, "pdf_images", return_value=observed):
            return VERIFY.verify(self.docx, self.pdf, self.update, self.evidence)

    def test_all_expected_screenshots_pass(self):
        self.save_docx(self.sources)
        audit = self.check(self.observed(self.sources))
        self.assertEqual(len(audit["matched"]), 2)

    def test_missing_word_screenshot_fails(self):
        self.save_docx(self.sources[:1])
        with self.assertRaisesRegex(ValueError, "omitted one or more embedded screenshots"):
            self.check(self.observed(self.sources))

    def test_missing_pdf_screenshot_fails(self):
        self.save_docx(self.sources)
        with self.assertRaisesRegex(ValueError, "rendered guide PDF omitted screenshots"):
            self.check(self.observed(self.sources[:1]))

    def test_unrelated_same_sized_pdf_image_fails(self):
        self.save_docx(self.sources)
        unrelated = self.root / "unrelated.png"
        Image.new("RGB", (80, 40), "green").save(unrelated)
        with self.assertRaisesRegex(ValueError, "rendered guide PDF omitted screenshots"):
            self.check(self.observed([self.sources[0], unrelated]))


if __name__ == "__main__":
    unittest.main()
