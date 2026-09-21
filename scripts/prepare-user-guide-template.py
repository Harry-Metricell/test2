"""Add the durable, hidden insertion marker required by the user-guide updater."""

from pathlib import Path
import argparse

from docx import Document


MARKER = "[[AUTO_GUIDE_CONTENT]]"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    args = parser.parse_args()
    template = Path(args.template)
    if not template.is_file():
        raise SystemExit(f"template not found: {template}")

    document = Document(str(template))
    if any(paragraph.text.strip() == MARKER for paragraph in document.paragraphs):
        print("marker already present")
        return

    # The marker is deliberately hidden: it remains a stable, own-paragraph
    # insertion point without becoming visible in a guide issued to users.
    paragraph = document.add_paragraph()
    run = paragraph.add_run(MARKER)
    run.font.hidden = True
    document.save(str(template))
    print("marker added")


if __name__ == "__main__":
    main()
