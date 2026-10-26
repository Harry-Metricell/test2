"""Regression tests for matching review criteria to tester screenshot evidence."""
import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from docx import Document
from PIL import Image


sys.dont_write_bytecode = True
MODULE = Path(__file__).with_name("build-evidence-report.py")
SPEC = importlib.util.spec_from_file_location("build_evidence_report", MODULE)
REPORT = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(REPORT)


class ResultMatchingTests(unittest.TestCase):
    def test_duplicate_assessment_is_removed_but_recovery_note_is_preserved(self):
        item = {'outcome': 'Passed', 'reason': 'The table shows the expected rows.'}
        result = {'reason': item['reason']}
        actual = item['reason'] + '\nEvidence: E01'
        self.assertEqual(REPORT.additional_criterion_assessment(item, result, actual), '')
        result['steps_taken'] = ['Recovered a tool timeout; cleanup completed.']
        extra = REPORT.additional_criterion_assessment(item, result, actual)
        self.assertIn('Recovered a tool timeout', extra)
        self.assertNotIn('The table shows the expected rows.', extra)
        self.assertEqual(REPORT.additional_criterion_assessment(item, {}, 'Different tester observation.\nEvidence: E01'), item['reason'])

    def test_review_cited_selected_capture_after_final_is_not_lost(self):
        images = [Path('criterion-3-' + suffix + '.png') for suffix in (
            'before-automobile', 'actual-automobile-config', 'before-add-service', 'final', 'final-selected', 'cleanup')]
        reason = 'criterion-3-before-add-service.png shows selection; criterion-3-final.png shows markers; criterion-3-final-selected.png shows selected result details.'
        selected = REPORT.reviewed_report_exhibits(images, 'Matching results are shown as selectable map points.', reason)
        self.assertEqual(selected, images[2:5])

    def test_sorting_keeps_reviewed_before_after_not_point_opening_setup(self):
        images = [Path('criterion-5-' + suffix + '.png') for suffix in ('before-select-point', 'before-sort', 'final')]
        reason = 'criterion-5-before-sort.png shows descending rows; criterion-5-final.png shows ascending rows.'
        self.assertEqual(REPORT.reviewed_report_exhibits(images, 'Sorting a supported table column reorders displayed rows.', reason), images[1:])

    def test_unknown_or_required_setup_is_not_pruned(self):
        images = [Path('criterion-1-' + suffix + '.png') for suffix in ('before-automobile', 'automobile-config', 'final')]
        self.assertEqual(REPORT.reviewed_report_exhibits(images, 'The layer selector includes Service.', 'criterion-1-final.png shows Service.'), images)
        self.assertEqual(REPORT.reviewed_report_exhibits(images, 'Results are displayed.', ''), REPORT.report_exhibits(images, 'Results are displayed.'))
        self.assertEqual(REPORT.reviewed_report_exhibits(images, 'The original map remains unchanged.', 'criterion-1-final.png shows the map.'), REPORT.report_exhibits(images, 'The original map remains unchanged.'))

    def test_table_enlargement_uses_measured_change_and_keeps_broad_maps_uncropped(self):
        with tempfile.TemporaryDirectory() as temporary:
            before = Path(temporary) / 'criterion-5-before-sort.png'
            final = Path(temporary) / 'criterion-5-final.png'
            Image.new('RGB', (1440, 900), 'white').save(before)
            changed = Image.new('RGB', (1440, 900), 'white')
            changed.paste('navy', (100, 550, 1200, 800))
            changed.save(final)
            self.assertIsNone(REPORT.localized_change_region([before, final]))
            self.assertIsNotNone(REPORT.criterion_focus_region([before, final], 'Sorting the table changes row order.', 'Passed'))
            Image.new('RGB', (1440, 900), 'blue').save(final)
            self.assertIsNone(REPORT.criterion_focus_region([before, final], 'Filtering the table changes rows.', 'Passed'))

    def test_passed_simulation_display_is_brief_and_preserves_source_audit(self):
        reason = ('SIMULATED UI STATE: The original request for pointId 123456 and mercatorKey 987654 was held and released. '
                  'criterion-8-baseline.png establishes the baseline; criterion-8-final.png visibly shows a small blue loading spinner. '
                  'criterion-8-recovery.png shows genuine details restored. Handler removed and pending work settled.')
        item = {'outcome': 'Passed', 'reason': reason}
        result = {'reason': reason, 'steps_taken': []}
        displayed = REPORT.criterion_assessment(item, result)
        self.assertIn('SIMULATED UI STATE:', displayed)
        self.assertIn('Delayed the selected detail request', displayed)
        self.assertIn('loading spinner', displayed)
        self.assertNotIn('123456', displayed)
        self.assertNotIn('establishes the baseline', displayed)
        self.assertLess(len(displayed), 240)
        self.assertEqual(item['reason'], reason)

    def test_failed_or_uncertain_simulation_keeps_its_full_explanation(self):
        reason = 'SIMULATED UI STATE: A request was held. criterion-4-final.png shows no spinner. Cleanup unverified.'
        self.assertEqual(REPORT.concise_simulation_assessment({'outcome': 'Failed', 'reason': reason}, {}), reason)
        passed = {'outcome': 'Passed', 'reason': reason}
        self.assertIn('Cleanup unverified', REPORT.concise_simulation_assessment(passed, {}))

    def test_image_cells_reserve_height_and_do_not_repeat_simulation_method(self):
        with tempfile.TemporaryDirectory() as temporary:
            image = Path(temporary) / 'criterion-8-final.png'
            Image.new('RGB', (1440, 900), 'white').save(image)
            document = Document()
            table = document.add_table(rows=1, cols=5)
            REPORT.add_image_row(table, [('E01', image, 'https://example.test/gis')])
            self.assertNotIn('SIMULATED UI STATE', table.rows[1].cells[0].text)
            nested = table.rows[1].cells[0].tables[0]
            picture = next(p for p in nested.cell(0, 0).paragraphs if p._p.xpath('.//w:drawing'))
            self.assertEqual(picture.paragraph_format.line_spacing, 1.0)
            self.assertTrue(nested.cell(0, 0)._tc.xpath('./w:tcPr/w:tcMar/w:bottom'))
            self.assertTrue(nested.rows[0]._tr.xpath('./w:trPr/w:cantSplit'))

    def test_all_passed_opening_limitations_have_no_column_bookkeeping(self):
        result = {'browserUrl': 'https://example.test/gis', 'browserVersion': '123'}
        displayed = REPORT.compact_limitations({}, [{'outcome': 'Passed'}], [result])
        self.assertEqual(displayed, 'No material testing limitations recorded.')
        self.assertNotIn('Y/N', displayed)

    def test_point_detail_exhibits_omit_layer_setup_but_keep_record_and_radio_proof(self):
        images = [Path('criterion-1-' + name + '.png') for name in (
            'before-surveyor', 'service-config', 'service-added', 'cluster',
            'selected-2g-3g-sa', 'final')]
        selected = REPORT.report_exhibits(images, 'The Service point details show serving-cell information.')
        self.assertEqual(selected, images[3:])
        configured = REPORT.report_exhibits(images, 'The Service point configuration settings remain unchanged.')
        self.assertIn(images[1], configured)

    def test_building_detail_exhibits_keep_selected_record_not_layer_setup(self):
        images = [Path('criterion-1-' + name + '.png') for name in (
            'building-config', 'building-map', 'single-building-map', 'after-point-selection', 'final')]
        selected = REPORT.report_exhibits(images, 'Selecting a Building point opens its separate dashboard.')
        self.assertEqual(selected, images[2:])

    def test_rebuild_preserves_exact_attempt_date_not_render_date_or_other_history(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / 'history').mkdir()
            results = [{'criterion': 1, 'outcome': 'Passed'}]
            self.assertEqual(REPORT.recorded_test_date(root / 'results.json', results), 'Not recorded')
            (root / 'history/attempt-001-test.json').write_text(json.dumps({
                'results': results, 'recordedAt': '2026-10-20T15:13:24.570Z'}))
            (root / 'history/attempt-002-test.json').write_text(json.dumps({
                'results': [], 'recordedAt': '2026-10-21T10:00:00Z'}))
            self.assertEqual(REPORT.recorded_test_date(root / 'results.json', results), '20/10/2026')
    def test_simulated_state_pairs_final_with_recovery_and_retains_record_baselines(self):
        images = [Path('criterion-6-' + name + '.png') for name in (
            'initial', 'gis-before-layer', 'genuine-target-baseline', 'distinct-valid-point',
            'final', 'recovery', 'cleanup')]
        selected = REPORT.report_exhibits(images, 'When point details fail to load, show an error.',
                                          'SIMULATED UI STATE: one scoped HTTP 500, then genuine recovery.')
        self.assertEqual(selected, [images[4], images[5], images[2], images[3]])
        # No explicit review disclosure means no shortcut or post-final reuse.
        self.assertNotIn(images[5], REPORT.report_exhibits(images, 'When point details fail to load, show an error.'))

    def test_point_content_keeps_identity_and_only_technology_configuration(self):
        images = [Path('criterion-3-' + name + '.png') for name in (
            'initial', 'surveyor-before-choice', 'tech-toggle-1', 'nsa-only-before-add',
            'cluster-before-point', 'final', 'cleanup')]
        selected = REPORT.report_exhibits(images, 'With 5G NSA as the only selected technology, opening a Ping point displays details.')
        self.assertEqual(selected, [images[3], images[5], images[4]])
        # An explicit toggle/restore journey must not be shortened to its end.
        selected = REPORT.report_exhibits(images, 'Toggle the point details twice and restore the original.')
        self.assertIn(images[2], selected)

    def test_opening_limitations_are_short_but_criterion_notes_keep_blockers_and_recovery(self):
        result = {'outcome': 'Blocked', 'reason': 'No available fixture. ' * 100,
                  'browserVersion': '123', 'browserUrl': 'https://example.test/gis'}
        outcome = {'outcome': 'Blocked', 'reason': 'A documented empty response is required.'}
        opening = REPORT.compact_limitations({}, [outcome], [result])
        self.assertIn('Blocked criterion(s): 1', opening)
        self.assertLess(len(opening), 300)
        self.assertIn(outcome['reason'], REPORT.criterion_assessment(outcome, result))
        recovered = {**result, 'outcome': 'Passed', 'reason': 'Recovered a tool timeout; cleanup completed.'}
        self.assertIn('Recovered a tool timeout', REPORT.criterion_assessment({'reason': 'State visibly correct.'}, recovered))

    def test_passed_simulation_remains_in_limitations_without_reviewer_disclosure(self):
        results = [{"outcome": "Passed", "reason": "SIMULATED UI STATE: documented empty response.",
                    "browserUrl": "https://example.test/gis", "browserVersion": "123"}]
        limitations = REPORT.report_limitations({}, [{"outcome": "Passed"}], results)
        self.assertIn("SIMULATED UI STATE: criterion(s) 1", limitations)
        self.assertIn("does not verify backend integration", limitations)
        self.assertNotIn("No testing limitations", limitations)

    def test_simulation_in_steps_is_disclosed_but_real_checks_are_not_labelled(self):
        base = {"outcome": "Passed", "browserUrl": "https://example.test/gis", "browserVersion": "123"}
        self.assertNotIn("SIMULATED UI STATE:", REPORT.report_limitations({}, [], [base]))
        simulated = {**base, "steps_taken": ["SIMULATED UI STATE: held original read request."]}
        self.assertIn("criterion(s) 2", REPORT.report_limitations({}, [], [base, simulated]))

    def test_generated_table_text_has_explicit_pdf_safe_font_without_changing_comparisons(self):
        document = Document()
        document.styles['Normal'].font.name = 'Calibri'
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        value = 'Criterion 7: All three MOS legends show 4 ≤ MOS ≤ 5. ' * 120
        REPORT.set_cell(cell, value)
        self.assertEqual(cell.text, value)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                self.assertEqual(run.font.name, 'Arial')
                self.assertEqual(run.font.size.pt, 9)
                self.assertEqual(run._r.rPr.rFonts.get(REPORT.qn('w:ascii')), 'Arial')
                self.assertEqual(run._r.rPr.rFonts.get(REPORT.qn('w:hAnsi')), 'Arial')

    def test_remains_same_pairs_each_recorded_state_and_search_state(self):
        images = [Path(name) for name in (
            "criterion-3-baseline-map.png", "criterion-3-baseline-layers.png",
            "criterion-3-before-search.png", "criterion-3-search-map-before-clear.png",
            "criterion-3-search-layers-before-clear.png", "criterion-3-final.png")]
        for role, baseline in (("map", images[0]), ("active-layer list", images[1])):
            with self.subTest(role=role):
                selected = REPORT.report_exhibits(images,
                    f"The visible {role} remains the same as its recorded initial state while searching and after the query is cleared.")
                self.assertEqual(selected[:2], [baseline, images[-1]])
                self.assertIn(images[3], selected)
                self.assertIn(images[4], selected)

    def test_clearing_restores_catalogue_pairs_original_and_final(self):
        images = [Path(name) for name in (
            "criterion-2-before-search.png", "criterion-2-before-clear.png", "criterion-2-final.png")]
        self.assertEqual(REPORT.report_exhibits(images,
            "Clearing the search restores the original catalogue controls and position."), [images[0], images[-1]])

    def test_single_unsaved_username_pairs_edit_and_reopened_empty_field(self):
        images = [Path(name) for name in (
            "criterion-2-before-open.png", "criterion-2-after-open.png",
            "criterion-2-after-edit-before-cancel.png", "criterion-2-after-cancel.png",
            "criterion-2-after-reopen.png", "criterion-2-final.png")]
        claim = "Enter an unsaved Username, cancel without adding, and reopen with an empty Username field."
        self.assertEqual(REPORT.report_exhibits(images, claim), [images[2], images[-1]])
        multiple = [*images[:3], Path("criterion-2-second-after-edit-before-cancel.png"), *images[3:]]
        self.assertIn(multiple[2], REPORT.report_exhibits(multiple, claim))
        self.assertIn(multiple[3], REPORT.report_exhibits(multiple, claim))
        unknown = [Path("step-a.png"), Path("step-b.png"), Path("criterion-2-final.png")]
        self.assertEqual(REPORT.report_exhibits(unknown, claim), unknown)
        self.assertNotEqual(REPORT.report_exhibits(images, claim.replace("empty", "not empty")), [images[2], images[-1]])

    def test_passed_restoration_does_not_enlarge_incidental_focus_changes(self):
        with tempfile.TemporaryDirectory() as temporary:
            before = Path(temporary) / "criterion-2-before-search.png"
            final = Path(temporary) / "criterion-2-final.png"
            Image.new("RGB", (1440, 900), "white").save(before)
            changed = Image.new("RGB", (1440, 900), "white")
            changed.paste("navy", (100, 100, 125, 125))
            changed.save(final)
            claim = "After the query is cleared, the original catalogue is restored."
            self.assertIsNotNone(REPORT.localized_change_region([before, final]))
            self.assertIsNone(REPORT.criterion_focus_region([before, final], claim, "Passed"))
            self.assertIsNone(REPORT.criterion_focus_region([before, final], "The map remains the same as its recorded initial state.", "Passed"))
            self.assertIsNotNone(REPORT.criterion_focus_region([before, final], claim, "Failed"))
            self.assertIsNotNone(REPORT.criterion_focus_region([before, final], "Changing the checkbox updates the setting.", "Passed"))

    def test_catalogue_clear_pairs_original_with_restored_view(self):
        images = [Path(name) for name in (
            "criterion-2-before-search.png", "criterion-2-after-search.png",
            "criterion-2-before-clear.png", "criterion-2-after-clear.png", "criterion-2-final.png")]
        self.assertEqual(REPORT.report_exhibits(images,
            "After the search query is cleared, the GIS layer catalogue shows the original contents at the same position."),
            [images[0], images[-1]])
        self.assertEqual(REPORT.report_exhibits(images,
            "The unmatched search shows no unrelated catalogue entries."), images)

    def test_unchanged_during_search_pairs_baseline_and_final_first(self):
        images = [Path(name) for name in (
            "criterion-4-baseline-map.png", "criterion-4-baseline-layers.png",
            "criterion-4-before-search.png", "criterion-4-after-search.png",
            "criterion-4-before-clear.png", "criterion-4-after-clear.png", "criterion-4-final.png")]
        self.assertEqual(REPORT.report_exhibits(images,
            "Searching and clearing the catalogue does not change the visible map state."),
            [images[0], images[-1], images[3]])
        self.assertEqual(REPORT.report_exhibits(images,
            "Searching and clearing the catalogue does not change the active layers."),
            [images[1], images[-1], images[3]])

    def test_recorded_initial_restoration_pairs_map_and_layers(self):
        images = [Path(name) for name in (
            "criterion-5-baseline-map.png", "criterion-5-baseline-layers.png",
            "criterion-5-before-search.png", "criterion-5-after-search.png", "criterion-5-final.png")]
        self.assertEqual(REPORT.report_exhibits(images,
            "After clearing the page is restored to its recorded initial catalogue, map, and active-layer state."),
            [images[0], images[-1], images[1], images[-1]])

    def test_multiple_layer_baselines_are_not_replaced_by_one_final(self):
        images = [Path(name) for name in (
            "criterion-3-baseline-layers-top.png", "criterion-3-baseline-layers-bottom.png",
            "criterion-3-after-search.png", "criterion-3-final.png")]
        self.assertEqual(REPORT.report_exhibits(images,
            "The complete original active layers remain unchanged."), images)
        combined = [images[0], images[1], Path("criterion-3-before-search.png"), *images[2:]]
        selected = REPORT.report_exhibits(combined,
            "After clearing the page is restored to its initial catalogue, map, and active layers.")
        self.assertIn(images[0], selected)
        self.assertIn(images[1], selected)

    def test_identical_search_restoration_endpoints_both_remain_visible(self):
        with tempfile.TemporaryDirectory() as temporary:
            before = Path(temporary) / "criterion-2-before-search.png"
            final = Path(temporary) / "criterion-2-final.png"
            Image.new("RGB", (80, 60), "navy").save(before)
            shutil.copyfile(before, final)
            self.assertEqual(REPORT.report_exhibits([before, final],
                "After the query is cleared, the original catalogue is restored."), [before, final])

    def test_map_and_layer_comparisons_are_paired_not_dialog_setup(self):
        images = [Path(name) for name in (
            "criterion-3-baseline-map.png", "criterion-3-baseline-layers.png",
            "criterion-3-first-open-before-edit.png", "criterion-3-first-edit-before-cancel.png",
            "criterion-3-final-reopened-empty-before-cancel.png",
            "criterion-3-final-restored-layers.png", "criterion-3-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "The map view and Map layers match the original pre-dialog baseline."), [images[0], images[-1], images[1], images[-2]])

    def test_routine_search_recovery_and_audit_chart_are_not_limitations(self):
        result = {"browserVersion": "1", "browserUrl": "https://example.test/gis",
                  "reason": "Saved the recovery transition and separate final PNG.",
                  "steps_taken": ["Waited for Requests and errors over time and observed the chart and 0.3% error rate."]}
        self.assertEqual(REPORT.report_limitations({}, [{"outcome": "Passed"}], [result]), "No testing limitations recorded.")
        result["steps_taken"].append("Recovered a selector error after a timeout.")
        self.assertIn("selector error", REPORT.report_limitations({}, [{"outcome": "Passed"}], [result]))
    def test_returned_launcher_is_navigation_baseline_not_earlier_audit(self):
        images = [Path(name) for name in (
            "criterion-4-initial.png", "criterion-4-after-open-audit-before-back.png",
            "criterion-4-after-back-before-open-gis.png", "criterion-4-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "From the returned launcher, open GIS and view the map."), images[2:])

    def test_filtered_configuration_omits_no_match_setup(self):
        images = [Path(name) for name in (
            "criterion-3-baseline-before-no-match.png", "criterion-3-after-no-match-before-replace.png",
            "criterion-3-after-replace-before-open.png", "criterion-3-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "With the filtered entry visible, open its configuration."), images[2:])

    def test_cancel_clear_keeps_original_catalogue_not_setup_search(self):
        images = [Path(name) for name in (
            "criterion-4-baseline-before-no-match.png", "criterion-4-after-no-match-before-replace.png",
            "criterion-4-after-replace-before-open.png", "criterion-4-after-open-before-cancel.png",
            "criterion-4-after-cancel-before-clear.png", "criterion-4-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "Cancel and clear Search layers; the original catalogue and map remain unchanged."), [images[0], *images[3:]])

    def test_open_dialog_does_not_replace_unchanged_map_baseline(self):
        images = [Path(name) for name in (
            "criterion-6-before-open-surveyor.png", "criterion-6-after-open-surveyor.png",
            "criterion-6-before-enter-first.png", "criterion-6-before-cancel-first.png",
            "criterion-6-before-enter-second.png", "criterion-6-before-cancel-second.png",
            "criterion-6-before-final-cancel.png", "criterion-6-final.png")]
        selected = REPORT.report_exhibits(images, "After cancelling the final reopened dialog, the original map view and Map layers entries/order are unchanged.")
        self.assertEqual(selected, [images[0], images[-1]])

    def test_second_edit_keeps_its_own_journey_not_first_edit_setup(self):
        images = [Path(name) for name in (
            "criterion-4-before-enter-first.png", "criterion-4-before-cancel-first.png",
            "criterion-4-before-reopen-first.png", "criterion-4-before-enter-second.png",
            "criterion-4-before-cancel-second.png", "criterion-4-before-reopen-second.png",
            "criterion-4-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "Reopening configuration after the second cancellation shows an empty Username field."), images[3:])
        self.assertEqual(REPORT.report_exhibits(images, "Reopening configuration after the second cancellation shows an empty Username field with both Cancel and Add layer available."), images[3:])
        self.assertEqual(REPORT.report_exhibits(images, "Both first and second cancellations discard their edits."), images)

    def test_identical_baseline_and_final_remain_a_visible_comparison(self):
        with tempfile.TemporaryDirectory() as temporary:
            baseline = Path(temporary) / "criterion-1-before-open-dialog.png"
            final = Path(temporary) / "criterion-1-final.png"
            Image.new("RGB", (80, 60), "navy").save(baseline)
            shutil.copyfile(baseline, final)
            self.assertEqual(REPORT.report_exhibits([baseline, final], "The original map remains unchanged."), [baseline, final])

    def test_single_opening_image_has_same_height_as_readable_paired_images(self):
        with tempfile.TemporaryDirectory() as temporary:
            image = Path(temporary) / "criterion-1-final.png"
            Image.new("RGB", (1440, 900), "navy").save(image)
            document = Document()
            table = document.add_table(rows=1, cols=5)
            REPORT.add_image_row(table, [("E01", image, "https://example.test/gis")], opening=True)
            self.assertAlmostEqual(document.inline_shapes[0].height.inches, 2.25)

    def test_report_text_removes_code_markers_without_changing_criteria(self):
        criterion = "Enter `Surveyor` and retain ``Map layers``."
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        REPORT.set_cell(cell, criterion)
        self.assertEqual(cell.text, "Enter Surveyor and retain Map layers.")
        self.assertEqual(REPORT.display_criterion(1, [criterion]), criterion)

    def test_search_baseline_excludes_configuration_and_keeps_full_journey(self):
        images = [Path(name) for name in (
            "criterion-2-after-open-gis.png", "criterion-2-after-open-surveyor-config.png",
            "criterion-2-before-confirm-add.png", "criterion-2-after-add-surveyor.png",
            "criterion-2-before-search.png", "criterion-2-before-search-layers-8.png",
            "criterion-2-after-search.png", "criterion-2-before-clear.png",
            "criterion-2-after-clear.png", "criterion-2-final.png",
            "criterion-2-cleanup.png",
        )]
        for criterion in (
            "After entering Surveyor in Search layers the list is unchanged.",
            "After clearing Search layers the original catalogue returns.",
            "After the search-and-clear journey no application-error page appears.",
        ):
            self.assertEqual(REPORT.report_exhibits(images, criterion), images[4:10])

    def test_setup_configuration_still_printed_when_it_is_the_assertion(self):
        images = [Path(name) for name in (
            "criterion-1-after-open-surveyor-config.png", "criterion-1-final.png",
        )]
        self.assertEqual(REPORT.report_exhibits(images, "The configuration dialog shows default values."), images)

    def test_negated_error_list_is_not_a_testing_limitation(self):
        result = {"browserVersion": "1", "browserUrl": "https://example.test/gis",
                  "reason": "No error or sign-in interruption.",
                  "steps_taken": ["GIS remained visible; no sign-in or application-error page appeared."]}
        self.assertEqual(REPORT.report_limitations({}, [{"outcome": "Passed"}], [result]), "No testing limitations recorded.")
        result["steps_taken"].append("Recovered a tool timeout without an error or sign-in interruption.")
        self.assertIn("Recovered a tool timeout", REPORT.report_limitations({}, [{"outcome": "Passed"}], [result]))

    def test_long_tables_repeat_column_header_and_keep_it_with_first_row(self):
        document = Document()
        table = document.add_table(rows=1, cols=5)
        REPORT.repeat_table_header(table)
        REPORT.repeat_table_header(table)
        self.assertEqual(len(table.rows[0]._tr.trPr.findall(REPORT.qn("w:tblHeader"))), 1)
        self.assertTrue(all(p.paragraph_format.keep_with_next for cell in table.rows[0].cells for p in cell.paragraphs))

    def test_reviewed_complete_recheck_replaces_incomplete_earlier_journey(self):
        images = [Path(name) for name in ("criterion-3-baseline-layers-0.png", "criterion-3-final.png",
                  "criterion-3-cleanup.png", "criterion-3-recheck-initial.png",
                  "criterion-3-recheck-baseline-layers-0.png", "criterion-3-recheck-baseline-layers-8.png",
                  "criterion-3-recheck-before-no-match.png", "criterion-3-recheck-asserted-layers-0.png",
                  "criterion-3-recheck-asserted-layers-8.png", "criterion-3-recheck-final.png", "criterion-3-recheck-cleanup.png")]
        selected = REPORT.report_exhibits(images, "The search leaves the Map layers list unchanged.",
                                         "The complete recheck supersedes earlier incomplete coverage.")
        self.assertEqual(selected, images[4:10])
        both = REPORT.report_exhibits(images, "The search leaves the Map layers list unchanged.")
        self.assertIn(images[0], both)
        self.assertIn(images[9], both)

    def test_search_keeps_assertion_not_unrelated_configuration(self):
        images = [Path(name) for name in ("criterion-2-after-open-surveyor-retry.png",
                  "criterion-2-after-add-defaults.png", "criterion-2-before-no-match.png",
                  "criterion-2-after-no-match.png", "criterion-2-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "After entering a value in Search layers, no entries are visible."), images[2:])

    def test_returned_launcher_audit_starts_at_own_launch_baseline(self):
        images = [Path(name) for name in ("criterion-3-initial.png", "criterion-3-before-back.png",
                  "criterion-3-after-back.png", "criterion-3-before-open-audit.png",
                  "criterion-3-after-open-audit.png", "criterion-3-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "Starting from the returned launcher, open API Request Audit."), images[3:])

    def test_default_configuration_with_before_add_name_is_not_lost(self):
        images = [Path(name) for name in ("criterion-1-after-open-gis.png", "criterion-1-after-open-surveyor-before-add.png", "criterion-1-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "A layer loaded with its default configuration is visible."), images[1:])

    def test_zoom_starts_at_its_recorded_map_baseline(self):
        images = [Path(name) for name in ("criterion-2-after-open-surveyor.png", "criterion-2-after-add.png", "criterion-2-settings-baseline.png", "criterion-2-before-zoom-in.png", "criterion-2-after-zoom-in.png", "criterion-2-before-zoom-out.png", "criterion-2-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "Zoom in and zoom out restore the original scale with the same layer."), images[3:])
        self.assertEqual(REPORT.report_exhibits(images, "After zooming, Show linked sites has its recorded checkbox state."), images[2:])

    def test_second_configuration_keeps_list_or_checkbox_baseline_not_first_add(self):
        images = [Path(name) for name in ("criterion-4-before-first-configuration.png", "criterion-4-after-first-configuration.png", "criterion-4-baseline-linked-sites-checked.png", "criterion-4-before-second-configuration.png", "criterion-4-before-cancel.png", "criterion-4-after-cancel.png", "criterion-4-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "Cancelling the second configuration leaves the Map layers list unchanged."), images[3:])
        self.assertEqual(REPORT.report_exhibits(images, "After cancelling the second configuration Show linked sites retains its checkbox state."), images[2:])

    def test_recovered_cancel_shows_decisive_pair_not_superseded_setup(self):
        images = [Path(name) for name in ("criterion-2-before-open.png", "criterion-2-final.png",
                  "criterion-2-before-open-recovered.png", "criterion-2-before-cancel-recovered.png",
                  "criterion-2-final-recovered.png", "criterion-2-cleanup.png")]
        self.assertEqual(REPORT.report_exhibits(images, "Select Cancel and verify the dialog closes."), images[3:5])

    def test_cancel_preserves_required_layer_list_baseline(self):
        images = [Path(name) for name in ("criterion-3-before-open.png", "criterion-3-before-cancel.png", "criterion-3-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "After cancelling the Map layers list matches the initial baseline."), images)

    def test_back_keeps_destination_and_return_not_launcher_setup(self):
        images = [Path(name) for name in ("criterion-3-initial.png", "criterion-3-after-open-before-back.png", "criterion-3-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "Using the browser Back control returns to the launcher."), images[1:])

    def test_only_direct_contradictions_count_as_defects(self):
        self.assertEqual(REPORT.confirmed_defect_count("Failed"), 1)
        for outcome in ("Passed", "Blocked", "Unverified"):
            self.assertEqual(REPORT.confirmed_defect_count(outcome), 0)

    def test_close_uses_actual_before_state_not_launcher(self):
        images = [Path(name) for name in ("criterion-3-initial.png", "criterion-3-before-add.png",
                  "criterion-3-after-add-dialog.png", "criterion-3-before-open.png",
                  "criterion-3-before-close.png", "criterion-3-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "When Display Settings is closed, the same layer remains."), images[4:])

    def test_defaults_preserve_configuration_and_unknown_transition_names(self):
        images = [Path(name) for name in ("criterion-1-initial.png", "criterion-1-before-add.png",
                  "criterion-1-after-add-dialog.png", "criterion-1-before-confirm-add.png",
                  "criterion-1-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "A layer loaded using its default configuration shows a legend."), images[3:])
        unknown = [Path(name) for name in ("criterion-2-initial.png", "criterion-2-before-expand.png", "criterion-2-final.png")]
        self.assertEqual(REPORT.report_exhibits(unknown, "The settings panel opens."), unknown[1:])

    def test_reopen_without_adding_does_not_include_add_setup(self):
        images = [Path(name) for name in ("criterion-4-initial.png", "criterion-4-before-add.png",
                  "criterion-4-after-add-dialog.png", "criterion-4-before-close.png",
                  "criterion-4-after-close.png", "criterion-4-final.png")]
        self.assertEqual(REPORT.report_exhibits(images, "Reopen the same layer without adding the layer again."), images[3:])

    def test_body_headings_visible_and_back_cover_preserved(self):
        document = Document(MODULE.parents[1] / "assets/templates/Automated Test Case Template.docx")
        REPORT.make_body_headings_visible(document)
        self.assertEqual(str(document.styles["Heading 1"].font.color.rgb), "000000")
        for paragraph in document.paragraphs:
            if paragraph.text.strip() == "Revision History":
                self.assertTrue(all(str(run.font.color.rgb) == "000000" for run in paragraph.runs))
            if paragraph.text.strip() == "About Metricell":
                self.assertTrue(all(str(run.font.color.rgb) == "FFFFFF" for run in paragraph.runs))

    def test_browser_details_are_observed_not_assumed(self):
        self.assertEqual(REPORT.browser_details([{}]), "Browser name/version not recorded")
        self.assertEqual(REPORT.browser_details([{"browserName": "Chromium", "browserVersion": "153.0.1.2"}]), "Chromium 153.0.1.2")

    def test_cleanup_only_cannot_be_used_as_fallback_evidence(self):
        with self.assertRaisesRegex(SystemExit, "only cleanup"):
            REPORT.report_exhibits([Path("criterion-1-cleanup.png")], "Layer is visible.")
        initial = [Path("criterion-1-initial.png")]
        self.assertEqual(REPORT.report_exhibits(initial, "Layer is visible."), initial)

    def test_continued_image_rows_keep_criterion_context(self):
        with tempfile.TemporaryDirectory() as temporary:
            image = Path(temporary) / "final.png"
            Image.new("RGB", (1440, 900), "navy").save(image)
            document = Document()
            table = document.add_table(rows=1, cols=5)
            REPORT.add_image_row(table, [("E01", image, "https://example.test/gis")], criterion_label="Criterion 4 continued: Reopen settings.")
            self.assertIn("Criterion 4 continued: Reopen settings.", table.rows[1].cells[0].text)

    def test_untested_block_builds_but_missing_claimed_evidence_fails(self):
        with tempfile.TemporaryDirectory(prefix="test2-no-browser-") as temporary:
            root = Path(temporary)
            template = MODULE.parent.parent / "assets/templates/Automated Test Case Template.docx"
            review = {"ticket": "TEST2-99", "overallOutcome": "Blocked", "criterionOutcomes": [{"criterion": "Launcher is visible.", "outcome": "Blocked", "reason": "Browser tools unavailable."}]}
            results = [{"criterion": "Launcher is visible.", "outcome": "Blocked", "evidence": [], "reason": "Browser tools unavailable."}]
            (root / "review.json").write_text(json.dumps(review))
            (root / "results.json").write_text(json.dumps(results))
            (root / "criteria.md").write_text("- [ ] Launcher is visible.\n")
            command = [sys.executable, str(MODULE), "--template", str(template), "--review-output", str(root / "review.json"), "--results", str(root / "results.json"), "--criteria", str(root / "criteria.md"), "--screenshots", str(root / "missing-folder"), "--output", str(root / "report.docx"), "--image-manifest", str(root / "manifest.json")]
            subprocess.run(command, check=True, capture_output=True)
            manifest = json.loads((root / "manifest.json").read_text())
            self.assertEqual(manifest["expectedEvidence"], [])
            self.assertIn("No browser evidence was captured", manifest["reportContent"]["limitations"])
            results[0]["evidence"] = ["missing.png"]
            (root / "results.json").write_text(json.dumps(results))
            failed = subprocess.run(command, capture_output=True, text=True)
            self.assertNotEqual(failed.returncode, 0)

    def test_uses_ordered_result_when_review_has_full_criterion_text(self):
        criteria = ["The launcher is displayed."]
        tester_result = {"criterion": 1, "evidence": ["criterion-1-initial.png", "criterion-1-final.png"]}
        reviewer_outcome = {"criterion": "The launcher is displayed."}
        matched = REPORT.result_for_outcome(reviewer_outcome, [tester_result], {"1": tester_result}, criteria, 0)
        self.assertEqual(matched, tester_result)

    def test_manifest_keeps_logical_references_but_deduplicates_identical_images(self):
        logical = [
            {"file": "criterion-1-initial.png", "sourceSha256": "same", "visualFingerprint": "a"},
            {"file": "criterion-2-final.png", "sourceSha256": "same", "visualFingerprint": "a"},
        ]
        expected = REPORT.unique_evidence_identities(logical)
        self.assertEqual(len(logical), 2)
        self.assertEqual(len(expected), 1)
        self.assertEqual(expected[0]["file"], "criterion-1-initial.png")

    def test_missing_referenced_screenshot_is_not_silently_omitted(self):
        with self.assertRaisesRegex(SystemExit, "referenced screenshot is missing"):
            REPORT.screenshot_paths({"evidence": ["missing.png"]}, [])

    def test_source_url_does_not_publish_query_secrets(self):
        self.assertEqual(REPORT.public_browser_url("https://user:secret@example.test/gis?token=secret#part"),
                         "https://example.test/gis")

    def test_report_uses_decisive_exhibits_not_repeated_setup(self):
        images = [Path(name) for name in (
            "criterion-4-initial.png", "criterion-4-after-open-gis.png",
            "criterion-4-before-restore.png", "criterion-4-after-restore.png",
            "criterion-4-before-close.png", "criterion-4-after-close.png",
            "criterion-4-after-reopen.png", "criterion-4-final.png", "criterion-4-after-cleanup.png")]
        selected = REPORT.report_exhibits(images, "Restore the control, close and reopen settings.")
        self.assertEqual([image.name for image in selected], [
            "criterion-4-before-restore.png", "criterion-4-after-restore.png",
            "criterion-4-before-close.png", "criterion-4-after-close.png",
            "criterion-4-final.png"])

    def test_final_is_retained_when_reopened_state_has_no_named_capture(self):
        images = [Path(name) for name in ("criterion-4-before-restore.png",
                  "criterion-4-after-restore.png", "criterion-4-after-close.png",
                  "criterion-4-final.png", "criterion-4-cleanup.png")]
        selected = REPORT.report_exhibits(images, "Restore, close and reopen settings.")
        self.assertIn(images[3], selected)
        self.assertNotIn(images[4], selected)

    def test_decisive_images_are_full_width_without_tiny_third_capture(self):
        with tempfile.TemporaryDirectory() as temporary:
            image = Path(temporary) / "final.png"
            Image.new("RGB", (1440, 900), "navy").save(image)
            document = Document()
            table = document.add_table(rows=1, cols=5)
            REPORT.add_image_row(table, [("E01", image, "https://example.test/gis")])
            self.assertGreater(document.inline_shapes[0].width.inches, 5)

    def test_back_cover_has_one_break_and_no_empty_spacer_page(self):
        document = Document()
        cases = document.add_table(rows=1, cols=5)
        document.add_page_break()
        document.add_paragraph()
        cover = document.add_paragraph("About Metricell")
        REPORT.remove_back_cover_spacers(cases)
        self.assertIs(cases._tbl.getnext(), cover._p)
        self.assertTrue(cover.paragraph_format.page_break_before)

    def test_cleanup_capture_is_not_a_decisive_exhibit(self):
        images = [Path(name) for name in (
            "criterion-4-before-close.png", "criterion-4-after-reopen.png",
            "criterion-4-final.png", "criterion-4-after-restore.png")]
        selected = REPORT.report_exhibits(images, "Restore, close and reopen settings.")
        self.assertNotIn(images[-1], selected)

    def test_report_reason_omits_file_presence_boilerplate(self):
        reason = ("Before and after screenshots show the change. The result browserUrl is "
                  "https://o2intelligence-v4-dev.metricell.com/gis, an expected V4 host. "
                  "All named criterion screenshots exist and are non-empty.")
        self.assertEqual(REPORT.concise_review_reason(reason), "Before and after screenshots show the change.")

    def test_steps_have_clean_separators(self):
        self.assertEqual(REPORT.formatted_steps(["Opened GIS.", "Changed the setting."]),
                         "Opened GIS; Changed the setting.")

    def test_setup_capture_selection_is_small_and_distinct(self):
        with tempfile.TemporaryDirectory(prefix="test2-setup-captures-") as temporary:
            folder = Path(temporary)
            files = [folder / name for name in (
                "criterion-1-after-open-gis.png", "criterion-2-after-open-gis.png",
                "criterion-1-before-load.png", "criterion-1-after-cleanup.png")]
            for index, file in enumerate(files):
                Image.new("RGB", (40, 30), (index * 40, 10, 10)).save(file)
            selected = REPORT.supporting_captures([(file, "https://example.test/gis") for file in files], {})
            self.assertEqual(len(selected), 2)
            self.assertEqual([item[1].name for item in selected], [files[0].name, files[1].name])

    def test_report_includes_outcome_limitations_urls_and_unique_evidence(self):
        with tempfile.TemporaryDirectory(prefix="test2-report-builder-") as temporary:
            root = Path(temporary)
            screenshots = root / "screenshots"
            screenshots.mkdir()
            first = screenshots / "criterion-1-initial.png"
            Image.new("RGB", (1440, 900), "blue").save(first)
            shutil.copy2(first, screenshots / "criterion-1-final.png")
            criteria = ["The launcher is displayed with all available module cards and each card has its own Open module control.", "A restricted user cannot open Surveyor."]
            (root / "criteria.md").write_text("".join(f"- [ ] {item}\n" for item in criteria), encoding="utf-8")
            (root / "results.json").write_text(json.dumps([
                {"criterion": 1, "outcome": "Unverified", "browserUrl": "https://example.test/launcher",
                 "reason": "Evidence is inconclusive.", "evidence": ["criterion-1-initial.png", "criterion-1-final.png"]},
                {"criterion": 2, "outcome": "Blocked", "reason": "No restricted account was supplied.", "evidence": []},
            ]), encoding="utf-8")
            (root / "review.json").write_text(json.dumps({
                "ticket": "TEST2-99", "overallOutcome": "Blocked", "reason": "Restricted account unavailable.",
                "criterionOutcomes": [
                    {"criterion": criteria[0], "outcome": "Unverified", "reason": "Evidence is inconclusive."},
                    {"criterion": criteria[1], "outcome": "Blocked", "reason": "No restricted account was supplied."},
                ],
            }), encoding="utf-8")
            output = root / "report.docx"
            manifest = root / "manifest.json"
            template = MODULE.parents[1] / "assets" / "templates" / "Automated Test Case Template.docx"
            subprocess.run([sys.executable, str(MODULE), "--template", str(template),
                            "--review-output", str(root / "review.json"), "--criteria", str(root / "criteria.md"),
                            "--results", str(root / "results.json"), "--screenshots", str(screenshots),
                            "--output", str(output), "--image-manifest", str(manifest)], check=True)
            document = Document(output)
            heading = next(paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "Test Cases")
            self.assertTrue(heading.paragraph_format.page_break_before)
            self.assertEqual(document.tables[3].cell(1, 0).text, f"1. {criteria[0]}")
            self.assertNotIn("...", document.tables[3].cell(1, 0).text)
            self.assertIsNotNone(document.tables[3].rows[1]._tr.trPr.find(REPORT.qn("w:cantSplit")))
            opening = heading._p
            while opening is not document.tables[5]._tbl:
                paragraphs = [opening] if opening.tag == REPORT.qn("w:p") else list(opening.iter(REPORT.qn("w:p")))
                for paragraph in paragraphs:
                    self.assertIsNotNone(paragraph.pPr.find(REPORT.qn("w:keepNext")))
                opening = opening.getnext()
                self.assertIsNotNone(opening)
            self.assertEqual(document.tables[5].cell(1, 0).text, "1.0.0")
            self.assertEqual(document.tables[5].cell(1, 1).text, criteria[0])
            cycle = document.tables[2]
            self.assertIn("Blocked criterion(s): 2", cycle.cell(5, 1).text)
            self.assertIn("Unverified criterion(s): 1", cycle.cell(5, 1).text)
            self.assertIn("not a confirmed application defect", cycle.cell(5, 1).text)
            self.assertEqual(document.tables[3].cell(1, 2).text, "Y")
            self.assertEqual(document.tables[3].cell(1, 1).text, "N")
            self.assertEqual(document.tables[3].cell(1, 3).text, "N")
            self.assertEqual(document.tables[3].cell(1, 4).text, "0")
            self.assertEqual(cycle.cell(6, 1).text, "Blocked")
            self.assertEqual(cycle.cell(7, 1).text, "https://example.test/launcher")
            self.assertEqual(cycle.cell(2, 1).text, "example.test")
            self.assertEqual(cycle.cell(3, 1).text, "Browser name/version not recorded")
            case_text = "\n".join(cell.text for row in document.tables[5].rows for cell in row.cells)
            self.assertIn("Evidence: E01", case_text)
            self.assertIn("Browser URL recorded after criterion: https://example.test/launcher", case_text)
            self.assertIn("Failed", case_text)  # Unverified remains Failed in the PDF convention.
            cases = document.tables[5]
            self.assertEqual(criteria[0], cases.rows[1].cells[1].text)
            image_row = next(row for row in cases.rows if row.cells[0].tables)
            self.assertIn("E01 - Final (criterion-1-final.png)", image_row.cells[0].tables[0].cell(0, 0).text)
            restricted_row = next(row for row in cases.rows if row.cells[0].text == '2.0.0')
            self.assertIn("A restricted user", restricted_row.cells[1].text)
            self.assertEqual([cell.text for cell in document.tables[3].rows[0].cells][1:],
                             ['Pass (Y/N)', 'Fail (Y/N)', 'Blocked (Y/N)', 'Defect No'])
            self.assertEqual([cell.text for cell in document.tables[3].rows[2].cells][1:],
                             ['N', 'N', 'Y', '0'])
            self.assertFalse(document.tables[3].autofit)
            self.assertGreater(document.tables[3].columns[0].width,
                               3 * document.tables[3].columns[1].width)
            self.assertNotIn('Criterion 1 assessment:', case_text)
            self.assertIn('Evidence is inconclusive.', cases.rows[1].cells[-2].text)
            self.assertFalse(cases.autofit)
            summary_heading = next(p for p in document.paragraphs if p.text.strip() == 'Test Summary')
            self.assertTrue(summary_heading.paragraph_format.page_break_before)
            self.assertNotIn("Supporting setup captures", case_text)
            images = json.loads(manifest.read_text(encoding="utf-8"))
            self.assertEqual(images["embeddedEvidenceImages"], 1)
            self.assertEqual(len(images["logicalEvidence"]), 1)
            self.assertEqual(images["reportContent"]["overallOutcome"], "Blocked")
            self.assertEqual(images["reportContent"]["evidenceCaptions"], ["E01 - Final (criterion-1-final.png)"])

    def test_adjacent_duplicate_final_is_printed_once(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            first, final = root / "after.png", root / "final.png"
            Image.new("RGB", (50, 30), "blue").save(first)
            shutil.copy2(first, final)
            self.assertEqual(REPORT.distinct_states([first, final]), [final])

    def test_one_changed_pixel_is_not_deduplicated(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            before, after = root / "before.png", root / "after.png"
            image = Image.new("RGB", (50, 30), "white")
            image.save(before)
            image.putpixel((25, 15), (0, 0, 0))
            image.save(after)
            self.assertEqual(REPORT.distinct_states([before, after]), [before, after])

    def test_return_to_original_state_is_not_deduplicated(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            before, changed, restored = [root / name for name in ("before.png", "changed.png", "restored.png")]
            Image.new("RGB", (50, 30), "blue").save(before)
            Image.new("RGB", (50, 30), "red").save(changed)
            shutil.copy2(before, restored)
            self.assertEqual(REPORT.distinct_states([before, changed, restored]), [before, changed, restored])

    def test_reopen_without_adding_does_not_request_setup(self):
        names = [Path(name) for name in ("criterion-3-after-add-surveyor.png", "criterion-3-before-confirm-defaults.png", "criterion-3-after-open-display-settings.png", "criterion-3-after-toggle.png", "criterion-3-after-close.png", "criterion-3-after-reopen.png")]
        self.assertEqual(REPORT.report_exhibits(names, "Reopen settings without adding the layer again."), names[2:])

    def test_setup_map_is_not_decisive_even_when_defaults_required(self):
        names = [Path(name) for name in ("criterion-1-after-open-gis.png", "criterion-1-before-add-surveyor.png", "criterion-1-before-confirm-defaults.png", "criterion-1-after-confirm-defaults.png", "criterion-1-before-open-display-settings.png", "criterion-1-final.png")]
        self.assertEqual(REPORT.report_exhibits(names, "Open settings on a layer with its default configuration."), names[2:])

    def test_passed_recovered_errors_are_recorded_without_changing_outcome(self):
        result = {"outcome": "Passed", "browserVersion": "1", "browserUrl": "https://example.test/gis", "reason": "Recovered add-layer tool timeout; cleanup removal reappeared and was completed."}
        notes = REPORT.report_limitations({}, [{"outcome": "Passed"}], [result])
        self.assertIn("Criterion 1: Recovered add-layer tool timeout", notes)
        self.assertNotIn("blocked", notes)
        self.assertEqual(result["outcome"], "Passed")

    def test_recovered_problem_can_be_found_in_steps(self):
        result = {"browserVersion": "1", "browserUrl": "https://example.test/gis", "reason": "Original state restored.", "steps_taken": ["Layer removal reappeared on return; removal was then verified."]}
        self.assertIn("removal reappeared", REPORT.report_limitations({}, [{"outcome": "Passed"}], [result]))

    def test_no_error_does_not_hide_real_recovered_timeout(self):
        result = {"browserVersion": "1", "browserUrl": "https://example.test/gis", "reason": "Recovered a tool timeout without application errors."}
        self.assertIn("Recovered a tool timeout", REPORT.report_limitations({}, [{"outcome": "Passed"}], [result]))
        result["reason"] = "Opened GIS without application errors."
        self.assertEqual(REPORT.report_limitations({}, [{"outcome": "Passed"}], [result]), "No testing limitations recorded.")

    def test_descriptive_caption_keeps_original_filename(self):
        self.assertEqual(REPORT.evidence_caption(Path("criterion-2-after-toggle.png")), "After toggle (criterion-2-after-toggle.png)")

    def test_localized_change_has_safe_enlargement_with_context(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            before, after = root / "criterion-1-before-toggle.png", root / "criterion-1-after-toggle.png"
            image = Image.new("RGB", (1440, 900), "white")
            image.save(before)
            from PIL import ImageDraw
            ImageDraw.Draw(image).rectangle((1100, 250, 1115, 265), fill="blue")
            image.save(after)
            box = REPORT.localized_change_region([before, after])
            self.assertIsNotNone(box)
            self.assertLess(box[0], 1100)
            self.assertGreater(box[2], 1115)
            self.assertLessEqual(box[2], 1440)
            document = Document()
            table = document.add_table(rows=1, cols=5)
            REPORT.add_image_row(table, [("E01", before, "https://example.test/gis"), ("E02", after, "https://example.test/gis")], focus_region=box)
            self.assertEqual(len(document.inline_shapes), 4)
            self.assertEqual(len(table._tbl.xpath('.//a:srcRect')), 2)
            self.assertIn("full screenshot above", table.rows[1].cells[0].tables[0].cell(1, 0).text)

    def test_broad_navigation_change_is_not_cropped(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            before, after = root / "criterion-1-before-open.png", root / "criterion-1-after-open.png"
            Image.new("RGB", (1440, 900), "white").save(before)
            Image.new("RGB", (1440, 900), "blue").save(after)
            self.assertIsNone(REPORT.localized_change_region([before, after]))

    def test_settings_exhibits_exclude_repeated_layer_setup(self):
        names = [Path(name) for name in (
            "criterion-1-after-open-layer-dialog.png", "criterion-1-after-add-layer.png",
            "criterion-1-before-toggle.png", "criterion-1-after-toggle.png", "criterion-1-final.png")]
        self.assertEqual(REPORT.report_exhibits(names, "Changing the control restores its original state."), names[2:])
        self.assertEqual(REPORT.report_exhibits(names[:2] + [names[-1]], "Opening Display Settings shows the control."), [names[-1]])

    def test_layer_setup_is_retained_when_it_is_the_assertion(self):
        names = [Path(name) for name in (
            "criterion-1-after-open-layer-dialog.png", "criterion-1-after-add-layer.png", "criterion-1-final.png")]
        self.assertEqual(REPORT.report_exhibits(names, "The layer is loaded with its default configuration."), names)
        self.assertEqual(REPORT.report_exhibits(names, "Opening the configuration dialog displays its controls."), names)
        self.assertEqual(REPORT.report_exhibits(names, "Adding a layer shows it in the layer list."), names[1:])

    def test_crop_is_bound_to_the_displayed_pair_not_setup(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            names = [root / name for name in (
                "criterion-1-after-open-layer-dialog.png", "criterion-1-after-add-layer.png",
                "criterion-1-before-toggle.png", "criterion-1-after-toggle.png", "criterion-1-final.png")]
            for file in names:
                Image.new("RGB", (1440, 900), "white").save(file)
            from PIL import ImageDraw
            changed = Image.open(names[3])
            ImageDraw.Draw(changed).rectangle((1100, 250, 1115, 265), fill="blue")
            changed.save(names[3])
            changed.save(names[4])
            self.assertIsNone(REPORT.localized_change_region(names[:2]))
            self.assertIsNotNone(REPORT.localized_change_region(names[2:4]))
            self.assertIsNotNone(REPORT.localized_change_region([names[2], names[4]]))
            self.assertIsNone(REPORT.localized_change_region([names[4]]))

    def test_restored_identical_state_is_printed_after_change(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            names = [root / name for name in ("criterion-1-before-toggle.png", "criterion-1-after-toggle.png", "criterion-1-final.png")]
            Image.new("RGB", (1440, 900), "white").save(names[0])
            Image.new("RGB", (1440, 900), "blue").save(names[1])
            shutil.copy2(names[0], names[2])
            (root / "criteria.md").write_text("- [ ] Restore the control.\n")
            (root / "results.json").write_text(json.dumps([{"criterion": 1, "outcome": "Passed", "evidence": [p.name for p in names], "browserUrl": "https://example.test/gis"}]))
            (root / "review.json").write_text(json.dumps({"ticket": "TEST2-99", "overallOutcome": "Passed", "criterionOutcomes": [{"criterion": "Restore the control.", "outcome": "Passed"}]}))
            output = root / "report.docx"
            subprocess.run([sys.executable, str(MODULE), "--template", str(MODULE.parents[1] / "assets/templates/Automated Test Case Template.docx"), "--review-output", str(root / "review.json"), "--criteria", str(root / "criteria.md"), "--results", str(root / "results.json"), "--screenshots", str(root), "--output", str(output), "--image-manifest", str(root / "manifest.json")], check=True)
            document = Document(output)
            self.assertEqual(len(document.inline_shapes), 3)
            self.assertEqual([cell.text for cell in document.tables[3].rows[1].cells][1:],
                             ['Y', 'N', 'N', '0'])
            manifest = json.loads((root / "manifest.json").read_text())
            self.assertEqual(manifest["embeddedEvidenceImages"], 2)
            self.assertEqual(len(manifest["logicalEvidence"]), 3)

    def test_reopen_baseline_uses_settled_before_reopen_capture(self):
        names = [Path(name) for name in ("criterion-3-before-toggle.png", "criterion-3-after-toggle.png", "criterion-3-after-close.png", "criterion-3-before-reopen.png", "criterion-3-after-reopen.png", "criterion-3-final.png")]
        self.assertEqual(REPORT.report_exhibits(names, "Reopening settings retains the changed state."), [names[0], names[1], names[3], names[5]])


if __name__ == "__main__":
    unittest.main()
