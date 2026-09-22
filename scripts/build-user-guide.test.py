"""Regression test for the deterministic, yellow-highlighted guide builder."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
BUILDER = ROOT / "scripts" / "build-user-guide.py"


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value), encoding="utf-8")


with tempfile.TemporaryDirectory() as temporary:
    repo = Path(temporary) / "repo"
    template = repo / "assets/user-guide/V4 User Guide Template.docx"
    template.parent.mkdir(parents=True)
    document = Document()
    document.add_heading("Example guide", 0)
    document.add_paragraph("Existing guidance must stay unchanged.")
    document.add_paragraph("[[AUTO_GUIDE_CONTENT]]")
    document.save(template)
    write_json(repo / "config/user-guide-plan.json", {"schema": "v4-user-guide-plan.v1", "baselineTemplate": "assets/user-guide/V4 User Guide Template.docx", "publishedGuide": "assets/user-guide/V4 User Guide.docx", "captureRoot": ".guide-staging", "updatePolicy": {"highlightAddedOrChangedContent": True, "highlightColour": "yellow", "preserveUnchangedContent": True, "requireScreenshotForChangedFeature": True}, "sections": []})
    write_json(repo / "config/user-guide-update-policy.json", {"schema": "v4-user-guide-update-policy.v1", "enabled": True, "onlyForPassedEvidenceReviews": True, "requireVerifiedTicketEvidence": True, "highlightColour": "yellow", "preserveExistingGuideContent": True, "allowChangeTypes": ["add", "amend"]})
    evidence = Path(temporary) / "evidence"
    image_path = evidence / "TEST2-99/screenshots/attempt-001/criterion-1-final.png"
    image_path.parent.mkdir(parents=True)
    Image.new("RGB", (64, 32), "navy").save(image_path)
    write_json(repo / "tickets/TEST2-99/guide-update.json", {"schema": "v4-user-guide-update.v1", "ticket": "TEST2-99", "title": "Open the launcher", "affectedSection": "Getting started", "changeType": "add", "steps": ["Open the launcher."], "screenshots": ["screenshots/attempt-001/criterion-1-final.png"], "reason": "The launcher is a new user-facing entry point."})
    command = [sys.executable, str(BUILDER), "--repo", str(repo), "--evidence-root", str(evidence), "--ticket", "TEST2-99"]
    first = subprocess.run(command, capture_output=True, text=True)
    assert first.returncode == 0, first.stderr or first.stdout
    output = repo / "assets/user-guide/V4 User Guide.docx"
    result = Document(output)
    text = "\n".join(paragraph.text for paragraph in result.paragraphs)
    assert "Existing guidance must stay unchanged." in text
    assert "New feature: Open the launcher" in text
    assert "[[AUTO_GUIDE_UPDATE:TEST2-99]]" in text
    assert text.count("[[AUTO_GUIDE_CONTENT]]") == 1
    assert "FFFF00" in result._element.xml, "new screenshot container must be yellow"
    second = subprocess.run(command, capture_output=True, text=True)
    assert second.returncode == 0, second.stderr or second.stdout
    repeat = Document(output)
    assert "\n".join(paragraph.text for paragraph in repeat.paragraphs).count("New feature: Open the launcher") == 1

print("build-user-guide regression test passed")
